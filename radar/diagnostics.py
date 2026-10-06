"""Offline bundled-runtime checks; never starts a search or opens user data."""
import io
import json
import ssl
import tempfile
from datetime import datetime
from pathlib import Path

from docx import Document
from lxml import html
from pypdf import PdfReader, PdfWriter
from .core import BERLIN
from .runtime import ROOT, FROZEN
from .store import Store


def self_test(report_path):
    checks = {}
    try:
        checks['frozen'] = FROZEN
        for name in ('index.html', 'app.js', 'style.css', 'favicon.svg'):
            assert (ROOT / 'web' / name).read_bytes()
        assert (ROOT / 'schedule.ps1').is_file()
        if FROZEN: assert (ROOT / 'THIRD_PARTY_NOTICES.txt').is_file()
        checks['assets'] = True
        assert datetime(2026, 1, 1, tzinfo=BERLIN).utcoffset().total_seconds() == 3600
        assert datetime(2026, 7, 1, tzinfo=BERLIN).utcoffset().total_seconds() == 7200
        checks['timezone'] = True
        assert ssl.create_default_context().verify_mode == ssl.CERT_REQUIRED
        checks['tls'] = True
        assert html.fromstring('<html><h1>Job Radar</h1></html>').xpath('//h1/text()') == ['Job Radar']
        checks['html_parser'] = True
        document = Document(); document.add_paragraph('Bundled CRM test')
        buffer = io.BytesIO(); document.save(buffer); buffer.seek(0)
        assert Document(buffer).paragraphs[0].text == 'Bundled CRM test'
        checks['docx'] = True
        writer = PdfWriter(); writer.add_blank_page(width=300, height=300)
        buffer = io.BytesIO(); writer.write(buffer); buffer.seek(0)
        assert len(PdfReader(buffer).pages) == 1
        checks['pdf'] = True
        with tempfile.TemporaryDirectory() as tmp:
            store = Store(Path(tmp) / 'test.sqlite')
            assert store.profile()['role_types']
        checks['sqlite'] = True
        if FROZEN:
            import ctypes
            assert ctypes.windll.advapi32.CredReadW
        checks['native_credentials_api'] = True
        result = {'ok': True, 'checks': checks}
    except Exception as exc:
        result = {'ok': False, 'checks': checks, 'error': str(exc)}
    report_path = Path(report_path)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(result, indent=2), encoding='utf-8')
    if not result['ok']: raise RuntimeError(result['error'])
