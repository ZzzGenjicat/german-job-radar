import io
import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path

from docx import Document
from reportlab.pdfgen import canvas

from radar.cv import save_and_extract_cv


def pdf_bytes(text):
    stream=io.BytesIO();pdf=canvas.Canvas(stream);pdf.drawString(72,720,text);pdf.save();return stream.getvalue()


def docx_bytes(text):
    stream=io.BytesIO();doc=Document();doc.add_paragraph(text);doc.save(stream);return stream.getvalue()


class CvTests(unittest.TestCase):
    def setUp(self):self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
    def tearDown(self):self.tmp.cleanup()

    def test_extracts_pdf_and_uses_generated_storage_name(self):
        result=save_and_extract_cv('../../Test Applicant CV.pdf',pdf_bytes('CRM Agent Python'),self.root)
        self.assertIn('CRM Agent Python',result['text'])
        self.assertEqual(result['kind'],'pdf')
        self.assertNotIn('Applicant',result['stored_name'])
        self.assertEqual((self.root/result['stored_name']).parent,self.root)

    def test_extracts_docx(self):
        result=save_and_extract_cv('Lebenslauf.docx',docx_bytes('KI Automatisierung und RAG'),self.root)
        self.assertIn('KI Automatisierung',result['text'])
        self.assertEqual(result['kind'],'docx')

    def test_extracts_docx_table_cells(self):
        stream=io.BytesIO();doc=Document();table=doc.add_table(rows=1,cols=2)
        table.cell(0,0).text='CRM Agent';table.cell(0,1).text='Python Automation';doc.save(stream)
        result=save_and_extract_cv('tabular-cv.docx',stream.getvalue(),self.root)
        self.assertIn('CRM Agent',result['text'])
        self.assertIn('Python Automation',result['text'])

    def test_rejects_mismatch_unsupported_and_oversize(self):
        with self.assertRaises(ValueError):save_and_extract_cv('resume.pdf',docx_bytes('wrong'),self.root)
        with self.assertRaises(ValueError):save_and_extract_cv('resume.exe',b'MZ',self.root)
        with self.assertRaises(ValueError):save_and_extract_cv('resume.pdf',b'%PDF-'+b'x'*(8*1024*1024),self.root)

    def test_failed_extraction_removes_untracked_original(self):
        with self.assertRaises(ValueError):
            save_and_extract_cv('empty.docx',docx_bytes(''),self.root)
        self.assertEqual(list(self.root.iterdir()),[])

    def test_parser_failure_returns_controlled_error_and_removes_file(self):
        with patch('radar.cv.PdfReader',side_effect=RuntimeError('parser failure')):
            with self.assertRaises(ValueError):save_and_extract_cv('bad.pdf',b'%PDF-invalid',self.root)
        self.assertEqual(list(self.root.iterdir()),[])


if __name__=='__main__':unittest.main()
